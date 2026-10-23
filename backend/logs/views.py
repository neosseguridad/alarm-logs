from django.shortcuts import render
from django.http import HttpResponse, JsonResponse
from django.views.decorators.csrf import csrf_exempt
from .models import Operator, AlarmType, Client, FalseAlarmType, AlarmLog
import json
from django.core.paginator import Paginator
from django.utils import timezone
import openpyxl
from openpyxl.utils import get_column_letter
from datetime import datetime

# Create your views here.
def get_operators(request):
    operators_list = []
    operators_queryset = Operator.objects.filter(is_active=True)
    for operator in operators_queryset:
        operators_list.append({
            "id" : operator.pk,
            "name" : f"{operator.last_name}, {operator.name}"
        })
    return HttpResponse(str(operators_list))

def get_alarm_types(request):
    alarm_types_list = []
    alarm_types_queryset = AlarmType.objects.filter(is_active=True)
    for alarm_type in alarm_types_queryset:
        alarm_types_list.append({
            "id" : alarm_type.pk,
            "name" : alarm_type.name
            })
    return HttpResponse(json.dumps(alarm_types_list))

def get_clients(request):
    clients_list = []
    clients_queryset = Client.objects.filter(is_active=True)
    for client in clients_queryset:
        clients_list.append({
            "id" : client.pk,
            "name" : client.name
        })
    return HttpResponse(str(clients_list))

def get_false_alarm_types(request):
    false_alarm_types_list = []
    false_alarm_types_queryset = FalseAlarmType.objects.filter(is_active=True)
    for false_alarm_type in false_alarm_types_queryset:
        false_alarm_types_list.append({
            "id" : false_alarm_type.pk,
            "name" : false_alarm_type.name})
    return HttpResponse(json.dumps(false_alarm_types_list))

def get_form_data(request):
    data = {
        "operators": [
            {
                "id": operator.pk,
                "name": f"{operator.last_name}, {operator.name}"
            }
            for operator in Operator.objects.filter(is_active=True)
        ],
        "alarm_types": [
            {
                "id": alarm_type.pk,
                "name": alarm_type.name
            }
            for alarm_type in AlarmType.objects.filter(is_active=True)
        ],
        "clients": [
            {
                "id": client.pk,
                "name": client.name
            }
            for client in Client.objects.filter(is_active=True)
        ],
        "false_alarm_types": [
            {
                "id": false_alarm_type.pk,
                "name": false_alarm_type.name
            }
            for false_alarm_type in FalseAlarmType.objects.filter(is_active=True)
        ]
    }

    return JsonResponse(data)

@csrf_exempt
def post_alarm_log(request):
    if request.method=="POST":
        data = json.loads(request.body)
        operator = Operator.objects.get(pk=data["operator"])
        client = Client.objects.get(pk=data["client"])
        alarm_type = AlarmType.objects.get(pk=data["alarm_type"])
        user_answered = data["user_answered"]
        alarm_type_desc = None
        if "alarm_type_desc" in data:
            alarm_type_desc = data["alarm_type_desc"]
        real_emergency = None
        if "real_emergency" in data:
            real_emergency = data["real_emergency"]
        false_alarm_type = None
        if "false_alarm_type" in data:
            false_alarm_type = FalseAlarmType.objects.get(pk=data["false_alarm_type"])
        false_alarm_type_desc = None
        if "false_alarm_type_desc" in data:
            false_alarm_type_desc = data["false_alarm_type_desc"]
        correct_password = None
        if "correct_password" in data:
            correct_password = data["correct_password"]
        correct_password_later = None
        if "correct_password_later" in data:
            correct_password_later = data["correct_password_later"]
        user_name = None
        if "user_name" in data:
            user_name = data["user_name"]
        log = AlarmLog.objects.create(
            is_active = True,
            operator = operator,
            client = client,
            user_name = user_name,
            alarm_type = alarm_type,
            user_answered = user_answered,
            alarm_type_desc = alarm_type_desc,
            real_emergency = real_emergency,
            false_alarm_type = false_alarm_type,
            false_alarm_type_desc = false_alarm_type_desc,
            correct_password = correct_password,
            correct_password_later = correct_password_later
        )
        return JsonResponse({
            "log_id" : log.pk
        })
    else:
        return(HttpResponse("Method not allowed"))

def get_log_info(request, id):
    try:
        log = AlarmLog.objects.get(pk=id, is_active=True)
        data = {
            "id" : log.pk,
            "datetime" : timezone.localtime(log.datetime).strftime("%Y/%m/%d %H:%M"),
            "operator" : f"{log.operator.last_name}, {log.operator.name}",
            "client" : log.client.name,
            "alarm_type" : log.alarm_type.name,
            "user_answered" : log.user_answered,
            "alarm_type_desc" : log.alarm_type_desc,
            "real_emergency" : log.real_emergency,
            "false_alarm_type" : log.false_alarm_type.name if log.false_alarm_type else None,
            "false_alarm_type_desc" : log.false_alarm_type_desc,
            "correct_password" : log.correct_password,
            "correct_password_later" : log.correct_password_later,
            "user_name" : log.user_name
        }
        return HttpResponse(json.dumps(data))
    except:
        return HttpResponse("not found")

def get_logs_page(request, page):
    queryset = AlarmLog.objects.filter(is_active=True)
    paginator = Paginator(queryset, 10)
    p = paginator.page(page)
    logs = []
    for log in p:
        logs.append({
            "id" : log.pk,
            "datetime" : timezone.localtime(log.datetime).strftime("%Y/%m/%d %H:%M"),
            "operator" : f"{log.operator.last_name}, {log.operator.name}",
            "client" : log.client.name
        })
    data = {
        "has_previous" : p.has_previous(),
        "has_next" : p.has_next(),
        "logs" : logs
    }

    return HttpResponse(json.dumps(data))

def get_logs_report(request):

    start_date = datetime.strptime(request.GET.get("start_date"), "%Y-%m-%d") if request.GET.get("start_date") else None
    finish_date = datetime.strptime(request.GET.get("finish_date"), "%Y-%m-%d") if request.GET.get("finish_date") else None

    queryset = AlarmLog.objects.filter(is_active=True)

    if start_date:
        queryset = queryset.filter(datetime__date__gte=start_date)
    if finish_date:
        queryset = queryset.filter(datetime__date__lte=finish_date)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alarm Logs"
    ws.append([
        "ID",
        "Fecha y Hora",
        "Operador",
        "Cliente",
        "Usuario contestó",
        "Nombre de usuario que contestó",
        "Tipo de Activación",
        "Descripción Tipo de Activación",
        "Emergencia Real",
        "Razón de Falsa Alarma",
        "Descripción Razón de Falsa Alarma",
        "Consigna Correcta",
        "Consigna Correcta Posteriormente",
    ])
    for log in queryset:
        ws.append([
            log.pk,
            timezone.localtime(log.datetime).strftime("%Y/%m/%d %H:%M"),
            f"{log.operator.last_name}, {log.operator.name}",
            log.client.name,
            "Sí" if log.user_answered else "No",
            log.user_name,
            log.alarm_type.name,
            log.alarm_type_desc,
            "Sí" if log.real_emergency else "No",
            log.false_alarm_type.name if log.false_alarm_type else None,
            log.false_alarm_type_desc,
            "Sí" if log.correct_password else "No",
            "Sí" if log.correct_password_later else "No"
        ])

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)

        for cell in col:
            if cell.value is not None:
                cell_len = len(str(cell.value))
                if cell_len > max_len:
                    max_len = cell_len
        ws.column_dimensions[col_letter].width = max_len + 3

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=Reporte_Alarmas.xlsx'
    wb.save(response)
    return response

